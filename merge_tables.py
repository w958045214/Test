#!/usr/bin/env python3
"""表格文件合并工具。

支持将多个 CSV / XLSX 文件按行合并，并输出为 CSV 或 XLSX。
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path
from typing import Iterable, List, Sequence, Tuple


class MergeError(Exception):
    """业务异常。"""


def _normalize_header(header: Sequence[str]) -> Tuple[str, ...]:
    return tuple((cell or "").strip() for cell in header)


def read_csv(path: Path) -> Tuple[Tuple[str, ...], List[List[str]]]:
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        try:
            header = _normalize_header(next(reader))
        except StopIteration as exc:
            raise MergeError(f"CSV 文件为空: {path}") from exc
        rows = [row for row in reader]
    return header, rows


def read_xlsx(path: Path) -> Tuple[Tuple[str, ...], List[List[str]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise MergeError(
            "读取 XLSX 需要安装 openpyxl，请先执行: pip install openpyxl"
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            raw_header = next(values)
        except StopIteration as exc:
            raise MergeError(f"XLSX 文件为空: {path}") from exc
        header = _normalize_header(["" if v is None else str(v) for v in raw_header])

        rows: List[List[str]] = []
        for raw_row in values:
            row = ["" if v is None else str(v) for v in raw_row]
            if len(row) < len(header):
                row.extend([""] * (len(header) - len(row)))
            rows.append(row)
    finally:
        workbook.close()

    return header, rows


def read_table(path: Path) -> Tuple[Tuple[str, ...], List[List[str]]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return read_csv(path)
    if suffix in {".xlsx", ".xlsm"}:
        return read_xlsx(path)
    raise MergeError(f"不支持的输入格式: {path.name}")


def write_csv(path: Path, header: Sequence[str], rows: Iterable[Sequence[str]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(header)
        writer.writerows(rows)


def write_xlsx(path: Path, header: Sequence[str], rows: Iterable[Sequence[str]]) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise MergeError(
            "输出 XLSX 需要安装 openpyxl，请先执行: pip install openpyxl"
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)


def write_table(path: Path, header: Sequence[str], rows: Iterable[Sequence[str]]) -> None:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        write_csv(path, header, rows)
        return
    if suffix == ".xlsx":
        write_xlsx(path, header, rows)
        return
    raise MergeError("输出文件扩展名仅支持 .csv 或 .xlsx")


def merge_tables(input_files: Sequence[Path], output_file: Path) -> int:
    if len(input_files) < 2:
        raise MergeError("至少需要提供 2 个输入文件")

    all_rows: List[List[str]] = []
    expected_header: Tuple[str, ...] | None = None

    for path in input_files:
        if not path.exists():
            raise MergeError(f"输入文件不存在: {path}")

        header, rows = read_table(path)
        if expected_header is None:
            expected_header = header
        elif header != expected_header:
            raise MergeError(
                "表头不一致，无法合并。\n"
                f"基准表头: {expected_header}\n"
                f"当前文件({path.name})表头: {header}"
            )

        for row in rows:
            if len(row) < len(expected_header):
                row.extend([""] * (len(expected_header) - len(row)))
            all_rows.append(row[: len(expected_header)])

    assert expected_header is not None
    write_table(output_file, expected_header, all_rows)
    return len(all_rows)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="合并多个表格文件（CSV/XLSX），按行拼接为一个输出文件。"
    )
    parser.add_argument(
        "-i",
        "--input",
        nargs="+",
        required=True,
        help="输入文件列表，至少 2 个（支持 .csv/.xlsx/.xlsm）",
    )
    parser.add_argument(
        "-o",
        "--output",
        required=True,
        help="输出文件路径（支持 .csv 或 .xlsx）",
    )
    return parser


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()

    input_files = [Path(p).expanduser().resolve() for p in args.input]
    output_file = Path(args.output).expanduser().resolve()

    try:
        merged_count = merge_tables(input_files, output_file)
    except MergeError as exc:
        parser.error(str(exc))

    print(f"合并完成，共写入 {merged_count} 行数据 -> {output_file}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
